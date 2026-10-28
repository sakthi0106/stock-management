"""
Comprehensive Backend Test Suite
"""
import unittest
import os
import io
import openpyxl
from datetime import datetime
from openpyxl.chart import BarChart, Reference
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app
from app.db.session import Base, get_db
from app.models.models import User, Product, StockTransaction, SyncRun
from app.core.security import hash_password, create_access_token, UserRole
from app.core.config import settings
from app.services.sync_service import SyncService
from app.services.excel_service import ExcelService

# Test database setup with StaticPool for in-memory SQLite shared across connections
TEST_DATABASE_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()

app.dependency_overrides[get_db] = override_get_db

class TestStockManagement(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        Base.metadata.create_all(bind=test_engine)
        cls.client = TestClient(app)
        
        # Seed test users
        db = TestingSessionLocal()
        cls.admin_user = User(
            email="test_admin@company.internal",
            hashed_password=hash_password("AdminSecret!"),
            full_name="Admin Test",
            role=UserRole.ADMIN,
            is_active=True
        )
        cls.director_user = User(
            email="test_director@company.internal",
            hashed_password=hash_password("DirectorSecret!"),
            full_name="Director Test",
            role=UserRole.DIRECTOR,
            is_active=True
        )
        cls.purchase_user = User(
            email="test_purchase@company.internal",
            hashed_password=hash_password("PurchaseSecret!"),
            full_name="Purchase Test",
            role=UserRole.PURCHASE_TEAM,
            is_active=True
        )
        db.add_all([cls.admin_user, cls.director_user, cls.purchase_user])
        db.commit()
        db.refresh(cls.admin_user)
        db.refresh(cls.director_user)
        db.refresh(cls.purchase_user)

        # Generate tokens
        cls.admin_token = create_access_token({"sub": str(cls.admin_user.id), "email": cls.admin_user.email, "role": cls.admin_user.role})
        cls.director_token = create_access_token({"sub": str(cls.director_user.id), "email": cls.director_user.email, "role": cls.director_user.role})
        cls.purchase_token = create_access_token({"sub": str(cls.purchase_user.id), "email": cls.purchase_user.email, "role": cls.purchase_user.role})
        db.close()

    def test_01_authentication_login(self):
        """Test valid and invalid login flows."""
        resp = self.client.post("/api/v1/auth/login", json={
            "email": "test_director@company.internal",
            "password": "DirectorSecret!"
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("access_token", data)
        self.assertEqual(data["role"], UserRole.DIRECTOR)

        bad_resp = self.client.post("/api/v1/auth/login", json={
            "email": "test_director@company.internal",
            "password": "WrongPassword!"
        })
        self.assertEqual(bad_resp.status_code, 401)

    def test_02_rbac_permissions(self):
        """Test role-based access control enforcement on endpoints."""
        # Purchase Team trying to access Admin-only user management route -> 403 Forbidden
        resp_forbidden = self.client.get(
            "/api/v1/admin/users",
            headers={"Authorization": f"Bearer {self.purchase_token}"}
        )
        self.assertEqual(resp_forbidden.status_code, 403)

        # Admin accessing user management -> 200 OK
        resp_admin = self.client.get(
            "/api/v1/admin/users",
            headers={"Authorization": f"Bearer {self.admin_token}"}
        )
        self.assertEqual(resp_admin.status_code, 200)
        users = resp_admin.json()
        self.assertGreaterEqual(len(users), 3)

    def test_03_excel_sync_and_deduplication(self):
        """Test Excel synchronization and verify deduplication idempotency."""
        sample_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
            "sample-data",
            "company_inventory_master.xlsx"
        )
        self.assertTrue(os.path.exists(sample_path), f"File not found: {sample_path}")
        with open(sample_path, "rb") as f:
            file_bytes = f.read()

        db = TestingSessionLocal()
        # First sync
        res1 = SyncService.execute_sync(
            db=db,
            file_bytes=file_bytes,
            file_name="company_inventory_master.xlsx",
            source_type="MANUAL_UPLOAD",
            triggered_by_user_id=self.admin_user.id
        )
        self.assertEqual(res1["status"], "SUCCESS")
        self.assertGreater(res1["rows_inserted"], 0)

        # Second sync of the EXACT SAME workbook
        res2 = SyncService.execute_sync(
            db=db,
            file_bytes=file_bytes,
            file_name="company_inventory_master.xlsx",
            source_type="MANUAL_UPLOAD",
            triggered_by_user_id=self.admin_user.id
        )
        self.assertEqual(res2["status"], "SUCCESS")
        # Idempotency check: No new rows inserted, all existing transactions skipped
        self.assertEqual(res2["rows_inserted"], 0)
        self.assertGreater(res2["rows_skipped"], 0)
        db.close()

    def test_04_dashboard_summary(self):
        """Test dashboard summary endpoint calculations."""
        resp = self.client.get(
            "/api/v1/dashboard/summary",
            headers={"Authorization": f"Bearer {self.director_token}"}
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("total_products", data)
        self.assertIn("total_stock_by_unit", data)
        self.assertIn("low_stock_count", data)
        self.assertIn("is_data_fresh", data)
        self.assertTrue(len(data["total_stock_by_unit"]) > 0)

    def test_05_products_and_movement_filtering(self):
        """Test product list, filtering by category, and movement history."""
        # Get products
        resp_prods = self.client.get(
            "/api/v1/products",
            headers={"Authorization": f"Bearer {self.purchase_token}"}
        )
        self.assertEqual(resp_prods.status_code, 200)
        prods = resp_prods.json()
        self.assertGreater(len(prods), 0)

        # Get movements
        resp_movs = self.client.get(
            "/api/v1/movements?transaction_type=INWARD",
            headers={"Authorization": f"Bearer {self.purchase_token}"}
        )
        self.assertEqual(resp_movs.status_code, 200)
        movs = resp_movs.json()
        for m in movs:
            self.assertEqual(m["transaction_type"], "INWARD")

    def test_06_reports_and_csv_export(self):
        """Test stock summary report and CSV export."""
        resp_csv = self.client.get(
            "/api/v1/reports/export?format=csv",
            headers={"Authorization": f"Bearer {self.director_token}"}
        )
        self.assertEqual(resp_csv.status_code, 200)
        self.assertIn("text/csv", resp_csv.headers["content-type"])
        csv_text = resp_csv.text
        self.assertIn("Product Code", csv_text)
        self.assertIn("PRD-1001", csv_text)

    def test_07_balance_after_tracking_and_case_insensitive_role(self):
        """Ensure transactions include running balances and role validation accepts lowercase values."""
        sample_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
            "sample-data",
            "company_inventory_master.xlsx"
        )
        with open(sample_path, "rb") as f:
            file_bytes = f.read()

        db = TestingSessionLocal()
        SyncService.execute_sync(
            db=db,
            file_bytes=file_bytes,
            file_name="company_inventory_master.xlsx",
            source_type="MANUAL_UPLOAD",
            triggered_by_user_id=self.admin_user.id
        )

        txn = db.query(StockTransaction).order_by(StockTransaction.id).first()
        self.assertIsNotNone(txn)
        self.assertIsNotNone(txn.balance_after)
        self.assertGreaterEqual(txn.balance_after, 0.0)
        db.close()

        resp = self.client.post(
            "/api/v1/admin/users",
            headers={"Authorization": f"Bearer {self.admin_token}"},
            json={
                "email": "lowercase-role-user@company.internal",
                "password": "StrongPass!123",
                "full_name": "Lowercase Role User",
                "role": "admin"
            }
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["role"], "ADMIN")

    def test_08_workbook_inspection_skips_chartsheets(self):
        """Ensure chart-only sheets do not break Excel inspection or enter sync mappings."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Stock Summary"
        worksheet.append(["Product Code", "Current Balance"])
        worksheet.append(["PART-001", 12])

        chart = BarChart()
        chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        workbook.create_chartsheet("Stock Chart").add_chart(chart)

        file_buffer = io.BytesIO()
        workbook.save(file_buffer)

        inspection = ExcelService.inspect_workbook(file_buffer.getvalue(), "inventory.xlsx")
        self.assertEqual([sheet["sheet_name"] for sheet in inspection["sheets"]], ["Stock Summary"])
        self.assertIn("Stock Summary", inspection["detected_recommendations"])
        self.assertNotIn("Stock Chart", inspection["detected_recommendations"])

    def test_09_daily_stock_register_import_and_resync(self):
        """Import row-2 material headers and keep dated daily movements current on re-sync."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Accessories"
        worksheet.cell(row=1, column=5, value="OPENING STOCK")
        worksheet.cell(row=1, column=7, value=datetime(2026, 10, 27))
        worksheet.cell(row=1, column=10, value=datetime(2026, 10, 28))
        headers = [
            "Description", "Part No", "Category", "UOM", "OPENING STOCK",
            "UNIT PRICE", "OPENING STOCK", "IN", "OUT", "OPENING STOCK", "IN", "OUT",
        ]
        for column, header in enumerate(headers, start=1):
            worksheet.cell(row=2, column=column, value=header)
        worksheet.append([
            "Test Washer", "DAILY-1001", "Fasteners", "PCS", 10, 2.5,
            8, 3, 1, 50, 20, 0,
        ])

        def workbook_bytes():
            buffer = io.BytesIO()
            workbook.save(buffer)
            return buffer.getvalue()

        original_bytes = workbook_bytes()
        inspection = ExcelService.inspect_workbook(original_bytes, "daily_inventory.xlsx")
        self.assertEqual(
            inspection["detected_recommendations"]["Accessories"]["likely_sheet_type"],
            "DAILY_STOCK_REGISTER"
        )

        db = TestingSessionLocal()
        first_sync = SyncService.execute_sync(
            db=db,
            file_bytes=original_bytes,
            file_name="daily_inventory.xlsx",
            source_type="MANUAL_UPLOAD",
        )
        self.assertEqual(first_sync["status"], "SUCCESS")
        self.assertEqual(first_sync["rows_read"], 1)

        product = db.query(Product).filter(Product.product_code == "DAILY-1001").one()
        self.assertEqual(product.opening_stock, 10)
        self.assertEqual(product.total_inward, 3)
        self.assertEqual(product.total_outward, 1)
        self.assertEqual(product.current_balance, 10)

        transactions = db.query(StockTransaction).filter(
            StockTransaction.product_id == product.id
        ).order_by(StockTransaction.transaction_type).all()
        self.assertEqual(len(transactions), 2)

        worksheet.cell(row=3, column=8, value=5)
        updated_sync = SyncService.execute_sync(
            db=db,
            file_bytes=workbook_bytes(),
            file_name="daily_inventory.xlsx",
            source_type="MANUAL_UPLOAD",
        )
        self.assertEqual(updated_sync["status"], "SUCCESS")
        self.assertEqual(product.current_balance, 12)
        self.assertEqual(product.total_inward, 5)
        self.assertEqual(
            db.query(StockTransaction).filter(StockTransaction.product_id == product.id).count(),
            2
        )
        inward = db.query(StockTransaction).filter(
            StockTransaction.product_id == product.id,
            StockTransaction.transaction_type == "INWARD",
        ).one()
        self.assertEqual(inward.quantity, 5)

        worksheet.cell(row=3, column=8).value = None
        cleared_sync = SyncService.execute_sync(
            db=db,
            file_bytes=workbook_bytes(),
            file_name="daily_inventory.xlsx",
            source_type="MANUAL_UPLOAD",
        )
        self.assertEqual(cleared_sync["status"], "SUCCESS")
        self.assertEqual(product.total_inward, 0)
        self.assertEqual(product.current_balance, 7)
        self.assertEqual(
            db.query(StockTransaction).filter(StockTransaction.product_id == product.id).count(),
            1
        )
        db.close()

    def test_10_agent_skips_identical_partial_import_with_rows(self):
        """A valid partial import should prevent the agent from needlessly re-importing the same file."""
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["Product Code", "Current Balance"])
        worksheet.append(["PARTIAL-1001", 10])
        file_buffer = io.BytesIO()
        workbook.save(file_buffer)
        file_bytes = file_buffer.getvalue()

        db = TestingSessionLocal()
        db.add(SyncRun(
            source_type="AGENT",
            status="PARTIAL",
            file_name="partial_inventory.xlsx",
            file_hash=ExcelService.calculate_file_hash(file_bytes),
            rows_read=1,
        ))
        db.commit()
        db.close()

        response = self.client.post(
            "/api/v1/sync/agent-push",
            headers={"X-Agent-Key": settings.SYNC_AGENT_API_KEY},
            files={"file": ("partial_inventory.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "UNCHANGED")

if __name__ == "__main__":
    unittest.main()
