"""
Stock Reports and Export Endpoints
Supports CSV and Excel generation for Directors, Purchase Teams, and Admins.
"""
import io
import csv
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import Response, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.db.session import get_db
from app.models.models import Product, StockTransaction, AuditLog
from app.core.security import require_roles, UserRole
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/reports", tags=["Reports"])

@router.get("/stock-summary")
def get_stock_report_data(
    category: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    current_user: dict = Depends(require_roles(UserRole.ALL)),
    db: Session = Depends(get_db)
):
    query = db.query(Product).filter(Product.is_active == True)
    if category:
        query = query.filter(Product.category == category)
        
    products = query.order_by(Product.product_code).all()
    results = []
    
    total_valuation = 0.0
    valuation_tracked_count = 0
    
    for p in products:
        st = "HEALTHY"
        if p.current_balance <= 0:
            st = "OUT_OF_STOCK"
        elif p.min_threshold > 0 and p.current_balance <= p.min_threshold:
            st = "LOW_STOCK"
            
        if status and status != "ALL" and st != status:
            continue
            
        row_val = None
        if p.unit_cost is not None and p.unit_cost > 0:
            row_val = round(p.current_balance * p.unit_cost, 2)
            total_valuation += row_val
            valuation_tracked_count += 1
            
        results.append({
            "product_code": p.product_code,
            "product_name": p.name,
            "category": p.category,
            "unit": p.unit_of_measure,
            "opening_stock": p.opening_stock,
            "total_inward": p.total_inward,
            "total_outward": p.total_outward,
            "current_balance": p.current_balance,
            "min_threshold": p.min_threshold,
            "unit_cost": p.unit_cost, # None if not provided in Excel
            "stock_value": row_val,    # Strictly computed only when real cost exists
            "bin_location": p.bin_location,
            "status": st
        })
        
    return {
        "items": results,
        "total_items": len(results),
        "total_valuation": round(total_valuation, 2) if valuation_tracked_count > 0 else None,
        "valuation_tracked_count": valuation_tracked_count
    }

@router.get("/export")
def export_stock_report(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    category: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    current_user: dict = Depends(require_roles(UserRole.ALL)),
    db: Session = Depends(get_db)
):
    query = db.query(Product).filter(Product.is_active == True)
    if category:
        query = query.filter(Product.category == category)
    products = query.order_by(Product.product_code).all()

    # Audit export event
    audit = AuditLog(
        user_id=int(current_user["sub"]),
        user_email=current_user.get("email"),
        action=f"EXPORT_STOCK_{format.upper()}",
        resource_type="REPORT",
        details=f"Exported {len(products)} products (filter: category={category}, status={status})"
    )
    db.add(audit)
    db.commit()

    fieldnames = [
        "Product Code", "Product Name", "Category", "Unit", 
        "Opening Stock", "Total Inward", "Total Outward", "Current Balance", 
        "Min Threshold", "Unit Price (INR)", "Stock Value (INR)", "Status", "Location"
    ]

    rows = []
    for p in products:
        st = "HEALTHY"
        if p.current_balance <= 0:
            st = "OUT_OF_STOCK"
        elif p.min_threshold > 0 and p.current_balance <= p.min_threshold:
            st = "LOW_STOCK"
        if status and status != "ALL" and st != status:
            continue

        val = round(p.current_balance * p.unit_cost, 2) if (p.unit_cost and p.unit_cost > 0) else ""
        rows.append([
            p.product_code,
            p.name,
            p.category,
            p.unit_of_measure,
            p.opening_stock,
            p.total_inward,
            p.total_outward,
            p.current_balance,
            p.min_threshold,
            p.unit_cost or "",
            val,
            st,
            p.bin_location or ""
        ])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(fieldnames)
        writer.writerows(rows)
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=Stock_Register_{timestamp}.csv"}
        )

    elif format == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock Register"
        ws.append(fieldnames)

        hdr_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        hdr_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(fieldnames) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = hdr_fill
            c.font = hdr_font

        for row in rows:
            ws.append(row)

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        out_stream = io.BytesIO()
        wb.save(out_stream)
        out_stream.seek(0)

        return StreamingResponse(
            out_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=Stock_Register_{timestamp}.xlsx"}
        )
